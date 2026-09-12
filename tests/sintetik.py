# -*- coding: utf-8 -*-
"""Sun'iy namuna fayl bilan avtomatik tekshiruv (CI uchun).

tests/golden.py haqiqiy bank fayllarini ishlatadi, ular esa mijozning
moliyaviy ma'lumoti — ochiq repozitoriyaga qo'yib bo'lmaydi va shuning
uchun GitHub Actions'da ishlamaydi. Bu yerdagi fayl esa kod bilan
yasaladi: hech qanday haqiqiy ma'lumot yo'q, lekin tuzilishi haqiqiy
"Сведения о работе счета" hisobotiga mos — sarlavha 6-qatorda, ИНН
ustuni yo'q, nom hisob raqam katagining ichida, summalar matn
ko'rinishida ("1 000 000,00").

Ishlatilishi:
    python tests/sintetik.py
Xato bo'lsa 1 qaytaradi, ya'ni CI qizil bo'ladi.
"""
import os
import sys
import tempfile
from decimal import Decimal

LOYIHA = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, LOYIHA)

# Har bir qator bitta qoidani tekshiradi.
QATORLAR = [
    # (hisob raqam, nom, op, debet, kredit, to'lov izohi, kutilgan guruh)
    ("23510000300961686800", "TEST TERMINAL", 4, "", "1 000 000,00",
     "00634 HUMO (1234567 - B16 TEST)", "Терминал"),
    ("20208000100000000001", "TEST BANK", 4, "50 000,00", "",
     "начисленные %% за обслуживание", "банк хизмати"),
    ("23402000300100001010", "TEST SOLIQ", 1, "200 000,00", "",
     "жисмоний шахслар даромад солиги учун тулов", "солик даромад"),
    ("23402000300100001010", "TEST GAZNA", 4, "", "300 000,00",
     "Возвратить по предприятию ПНФЛ согласно заключения ГНИ № 1", "ДСИ"),
    ("20214000404954805001", "TEST AROMAT", 1, "70 000,00", "",
     "услуга ароматизации воздуха в помещениях", "ароматизация"),
    ("20208000500810057001", "NOMALUM MCHJ", 1, "400 000,00", "",
     "шартнома асосан тулов", "?"),
]


def namuna_yasash(yol):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сведения о работе счета"

    ws.cell(1, 1, "00083 / SINOV FILIALI")
    ws.cell(1, 6, "ABS/Клиент-Банк")
    ws.cell(2, 1, "Сведения о работе счета c 01.08.2026 по 31.08.2026")
    ws.cell(3, 1, "Счет: 20218000900961686001")
    ws.cell(3, 6, "SINOV TASHKILOTI")
    # Qoldiqlar aylanma bilan mos bo'lishi shart:
    #   1 000 + 1 300 000 (kredit) - 720 000 (debet) = 581 000
    # Balans nazorati aynan shu tenglikni tekshiradi.
    ws.cell(4, 1, "Остаток на начало периода: 1 000,00")
    ws.cell(4, 6, "Остаток на конец периода: 581 000,00")

    sarlavhalar = ["Дата", "Счет", "№ док", "Оп", "МФО",
                   "Оборот Дебет", "Оборот Кредит", "Назначение платежа"]
    for i, s in enumerate(sarlavhalar, start=1):
        ws.cell(6, i, s)

    for n, (hisob, nom, op, debet, kredit, izoh, _kutilgan) in enumerate(QATORLAR):
        r = 7 + n
        ws.cell(r, 1, "03.08.2026")
        # Haqiqiy hisobotda nom aynan shunday — hisob raqamdan keyin
        # yangi qatordan boshlab yoziladi.
        ws.cell(r, 2, f"{hisob}\n{nom}")
        ws.cell(r, 3, str(100000 + n))
        ws.cell(r, 4, op)
        ws.cell(r, 5, "00083")
        ws.cell(r, 6, debet or "0,00")
        ws.cell(r, 7, kredit or "0,00")
        ws.cell(r, 8, izoh)

    wb.save(yol)


# "Справка о работе счета" (D formati) uchun alohida namuna. Bu formatda
# ustunlar ikki tomonga bo'lingan — mijoz (bizning hisobimiz) va
# korrespondent (kontragent) — va ma'lumot 1-emas, 2-ustundan boshlanadi.
# Ikkita tuzoq bor, ikkalasi ham shu yerda tekshiriladi:
#   1) sarlavhalarda "ё" ishlatiladi ("Счёт корреспондента")
#   2) to'lov izohi katak ichida bir necha qatorga bo'lingan holda keladi
NL = chr(10)  # izoh ichidagi qator uzilishi (heredoc bilan yozib bo'lmaydi)
D_QATORLAR = [
    # (korrespondent nomi, ИНН, hisob, debet, kredit, izoh, kutilgan guruh)
    ("TEST TERMINAL", "300000001", "23510000300961686800", 0, 1000000,
     "0063400634 Возмещение клиенту по покупкам" + NL + "ТСП договор humo", "Терминал"),
    ("TEST QURIQLASH", "200237592", "20208000100000000002", 150000, 0,
     "куриклаш хизмати учун" + NL + "тулов", "куриклаш"),
    ("TEST IJTIMOIY", "300000003", "23402000300100001010", 250000, 0,
     "ижтимоий солик" + NL + "учун тулов", "солик ижтимоий"),
]


# Bir xil formatning ikki tili. Bank interfeysi qaysi tilda bo'lsa,
# hisobot ham o'sha tilda keladi — ustun nomlari boshqa, tuzilishi bir xil.
D_SARLAVHALAR = {
    "ru": {
        2: "№", 3: "№ Док-та", 4: "ВО", 5: "Дата",
        6: "Наименование клиента", 7: "ИНН клиента",
        8: "Счёт клиента", 9: "МФО клиента",
        10: "Сумма дебета", 11: "Сумма кредита",
        12: "Наименование корреспондента", 13: "ИНН корреспондента",
        14: "Счёт корреспондента", 15: "МФО корреспондента",
        16: "Назначение платежа",
        "bosh": "Остаток на начало периода Пассив",
        "oxir": "Остаток на конец периода Пассив",
    },
    "uz": {
        2: "№", 3: "№ Ҳуж-та", 4: "ВО", 5: "Сана",
        6: "Номи мижоз", 7: "СТИР мижоз",
        8: "Ҳисоб мижоз", 9: "МФО мижоз",
        10: "Дебет суммаси", 11: "Кредит суммаси",
        12: "Номи корреспондент", 13: "СТИР корреспондент",
        14: "Ҳисоб корреспондент", 15: "МФО корреспондент",
        16: "Тўлов мақсади",
        "bosh": "Давр бошидаги қолдиқ Пассив",
        "oxir": "Давр охиридаги қолдиқ Пассив",
    },
}


def namuna_d_yasash(yol, til="ru"):
    """D formatidagi namuna. Mijoz ustunlari ataylab boshqa qiymat bilan
    to'ldiriladi: kod kontragentni emas, bizning firmani o'qib qo'ysa,
    sinov darhol yiqiladi."""
    import datetime
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Справка о работе счета"

    ws.cell(2, 2, "00083 / SINOV FILIALI")
    ws.cell(3, 2, "История по счёту: 20218000900961686001")
    ws.cell(4, 2, "SINOV TASHKILOTI")
    ws.cell(5, 2, "Справка о работе счёта ЗА 01.09.2026 - 30.09.2026")

    debet = sum(q[3] for q in D_QATORLAR)
    kredit = sum(q[4] for q in D_QATORLAR)
    sarlavhalar = D_SARLAVHALAR[til]
    ws.cell(6, 2, sarlavhalar["bosh"])
    ws.cell(6, 6, 1000)

    for c, nom in sarlavhalar.items():
        if isinstance(c, int):
            ws.cell(7, c, nom)

    for n, (nomi, inn, hisob, deb, kred, izoh, _kut) in enumerate(D_QATORLAR):
        r = 8 + n
        ws.cell(r, 2, n + 1)
        ws.cell(r, 3, f"000004{n}")
        ws.cell(r, 4, "06")
        ws.cell(r, 5, datetime.datetime(2026, 9, 7 + n))
        ws.cell(r, 6, "BIZNING FIRMA")          # mijoz — o'qilmasligi kerak
        ws.cell(r, 7, "599656770")              # mijoz ИНН — o'qilmasligi kerak
        ws.cell(r, 8, "20218000900961686001")   # mijoz hisobi — o'qilmasligi kerak
        ws.cell(r, 9, "00083")
        ws.cell(r, 10, float(deb))
        ws.cell(r, 11, float(kred))
        ws.cell(r, 12, nomi)
        ws.cell(r, 13, inn)
        ws.cell(r, 14, hisob)
        ws.cell(r, 15, "01158")
        ws.cell(r, 16, izoh)

    oxir = 8 + len(D_QATORLAR)
    ws.cell(oxir, 2, sarlavhalar["oxir"])
    ws.cell(oxir, 6, 1000 + kredit - debet)
    wb.save(yol)


def d_formatini_tekshir(core, ish, til="ru"):
    """D formati o'qiladimi, kontragent ustunlari to'g'ri tanlanadimi."""
    xatolar = []
    xom = os.path.join(ish, f"sinov_d_{til}.xlsx")
    namuna_d_yasash(xom, til)

    wb, ws, layout, rows = core.load_raw_rows(xom)
    if len(rows) != len(D_QATORLAR):
        xatolar.append(f"D/{til}: {len(rows)} qator o'qildi, {len(D_QATORLAR)} kutilgan")
        return xatolar

    for r, (nomi, inn, hisob, _d, _k, _izoh, kutilgan) in zip(rows, D_QATORLAR):
        # Mijoz emas, korrespondent ustunlari olinganini tekshiramiz.
        if r["name"] != nomi:
            xatolar.append(f"D/{til}: nom '{r['name']}' o'qildi, '{nomi}' kutilgan")
        if str(r["inn"]) != inn:
            xatolar.append(f"D/{til}: ИНН '{r['inn']}' o'qildi, '{inn}' kutilgan")
        if str(r["account"]) != hisob:
            xatolar.append(f"D/{til}: hisob '{r['account']}' o'qildi, '{hisob}' kutilgan")
        guruh, _ = core.classify_row(r["op"], r["name"], r["purpose"], r["inn"], r["account"])
        if guruh != kutilgan:
            xatolar.append(f"D/{til}: {nomi} -> '{guruh}', '{kutilgan}' kutilgan edi")

    natija = os.path.join(ish, f"natija_d_{til}.xlsx")
    info = core.build_simplified_report(xom, natija)
    if info.get("balans_farqi") != Decimal(0):
        xatolar.append(f"D/{til}: balans farqi {info.get('balans_farqi')}")
    return xatolar


def main():
    import core

    ish = tempfile.mkdtemp()
    xom = os.path.join(ish, "sinov.xlsx")
    namuna_yasash(xom)

    xatolar = []

    # 1) Har bir qator kutilgan guruhga tushdimi
    wb, ws, layout, rows = core.load_raw_rows(xom)
    if len(rows) != len(QATORLAR):
        xatolar.append(f"qatorlar soni: {len(rows)}, kutilgani {len(QATORLAR)}")
    for r, (_h, nom, _op, _d, _k, _izoh, kutilgan) in zip(rows, QATORLAR):
        guruh, _ishonch = core.classify_row(
            r["op"], r["name"], r["purpose"], r["inn"], r["account"]
        )
        if guruh != kutilgan:
            xatolar.append(f"{nom}: '{guruh}' chiqdi, '{kutilgan}' kutilgan edi")

    # 2) Hisobot yasaladimi va yig'indi to'g'rimi
    natija = os.path.join(ish, "natija.xlsx")
    info = core.build_simplified_report(xom, natija)

    import openpyxl

    lst = openpyxl.load_workbook(natija, data_only=True)["Лист1"]
    kategoriya_debet = Decimal("0")
    kategoriya_kredit = Decimal("0")
    jami_qator = None
    for row in range(4, lst.max_row + 1):
        nom = lst.cell(row, 2).value
        if nom is None:
            continue
        d = lst.cell(row, 3).value or 0
        k = lst.cell(row, 4).value or 0
        if str(nom).strip() == "Общий итог":
            jami_qator = (Decimal(str(d)), Decimal(str(k)))
        else:
            kategoriya_debet += Decimal(str(d))
            kategoriya_kredit += Decimal(str(k))

    if jami_qator is None:
        xatolar.append("Лист1 da 'Общий итог' qatori topilmadi")
    else:
        # Yig'ma jadval o'z yig'indisiga teng bo'lishi shart: bu jimgina
        # buziladigan xato turi, ko'z bilan sezilmaydi.
        if jami_qator[0] != kategoriya_debet:
            xatolar.append(f"debet yig'indisi: {jami_qator[0]} != {kategoriya_debet}")
        if jami_qator[1] != kategoriya_kredit:
            xatolar.append(f"kredit yig'indisi: {jami_qator[1]} != {kategoriya_kredit}")
        if jami_qator[0] != info["total_debet"] or jami_qator[1] != info["total_kredit"]:
            xatolar.append("Лист1 yig'indisi hisoblangan jami bilan mos emas")

    # 3) Balans nazorati: bank ko'rsatgan qoldiqlar bilan aylanma mos kelsin
    if info.get("balans_farqi") != Decimal(0):
        xatolar.append(f"balans nazorati: farq {info.get('balans_farqi')}, 0 kutilgan edi")

    for til in ("ru", "uz"):
        xatolar += d_formatini_tekshir(core, ish, til)

    if xatolar:
        print("XATO:")
        for x in xatolar:
            print("  -", x)
        return 1

    print(f"[OK] sintetik sinov: {len(QATORLAR)} qator + D formati (ru, uz), "
          f"jami {jami_qator[0]} debet / {jami_qator[1]} kredit")
    return 0


if __name__ == "__main__":
    sys.exit(main())
