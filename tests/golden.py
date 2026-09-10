# -*- coding: utf-8 -*-
"""Etalon (golden) test — kod o'zgarganda raqamlar o'zgarmaganini tekshiradi.

Nima uchun kerak: kod GitHub'dan yuklanadi, ya'ni `git push` darhol barcha
foydalanuvchiga boradi. Launcher faqat kod YUKLANMASA zaxiraga qaytadi —
kod yuklanib, lekin noto'g'ri raqam chiqarsa hech kim sezmaydi.
Buxgalteriyada bu eng yomon xato turi: hisobot ishlagandek ko'rinadi.

Ishlatilishi:
    python tests/golden.py               # solishtiradi
    python tests/golden.py --yozib-ol    # joriy natijani etalon deb saqlaydi

Kirish fayllari: sinov/*.xlsx (haqiqiy bank hisobotlari).
Etalonlar:       tests/golden/*.json

Ikkalasi ham .gitignore da — ular mijozning moliyaviy ma'lumoti, ochiq
repozitoriyaga chiqmasligi kerak. Ya'ni bu test push'dan OLDIN mahalliy
ishga tushiriladi.
"""
import glob
import json
import os
import sys
import tempfile

LOYIHA = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, LOYIHA)

SINOV_DIR = os.path.join(LOYIHA, "sinov")
ETALON_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "golden")


def _xulosa(out_path):
    """Natija faylidan Лист1 xulosasini o'qiydi.

    Butun faylni emas, aynan xulosani solishtiramiz: uslub, ustun kengligi
    yoki xom varaqdagi tartib o'zgarishi testni yiqitmasligi kerak —
    faqat raqamlar va kategoriyalar muhim."""
    import openpyxl

    ws = openpyxl.load_workbook(out_path, data_only=True)["Лист1"]
    qatorlar = []
    for r in range(1, ws.max_row + 1):
        nom = ws.cell(r, 2).value
        if nom is None:
            continue
        debet, kredit = ws.cell(r, 3).value, ws.cell(r, 4).value
        qatorlar.append({
            "nom": str(nom).strip(),
            "debet": None if debet is None else str(debet),
            "kredit": None if kredit is None else str(kredit),
        })
    return qatorlar


def _natija(src):
    import core

    out = os.path.join(tempfile.mkdtemp(), "natija.xlsx")
    info = core.build_simplified_report(src, out)
    return {
        "review_count": info["review_count"],
        "total_debet": str(info["total_debet"]),
        "total_kredit": str(info["total_kredit"]),
        "xulosa": _xulosa(out),
    }


def _etalon_yoli(src):
    return os.path.join(ETALON_DIR, os.path.basename(src) + ".json")


def _farqlar(etalon, joriy):
    farq = []
    for kalit in ("review_count", "total_debet", "total_kredit"):
        if etalon[kalit] != joriy[kalit]:
            farq.append(f"  {kalit}: {etalon[kalit]} -> {joriy[kalit]}")

    eski = {q["nom"]: q for q in etalon["xulosa"]}
    yangi = {q["nom"]: q for q in joriy["xulosa"]}
    for nom in sorted(set(eski) | set(yangi)):
        if nom not in yangi:
            farq.append(f"  '{nom}' kategoriyasi YO'QOLDI")
        elif nom not in eski:
            farq.append(f"  '{nom}' kategoriyasi PAYDO BO'LDI: {yangi[nom]}")
        elif eski[nom] != yangi[nom]:
            farq.append(
                f"  '{nom}': debet {eski[nom]['debet']} -> {yangi[nom]['debet']}, "
                f"kredit {eski[nom]['kredit']} -> {yangi[nom]['kredit']}"
            )
    return farq


def main():
    yozib_ol = "--yozib-ol" in sys.argv
    fayllar = sorted(glob.glob(os.path.join(SINOV_DIR, "*.xlsx")))
    if not fayllar:
        print(f"sinov fayllari topilmadi: {SINOV_DIR}")
        return 1

    os.makedirs(ETALON_DIR, exist_ok=True)
    xato = 0
    for src in fayllar:
        nom = os.path.basename(src)
        joriy = _natija(src)
        yol = _etalon_yoli(src)

        if yozib_ol or not os.path.exists(yol):
            with open(yol, "w", encoding="utf-8") as f:
                json.dump(joriy, f, ensure_ascii=False, indent=2)
            holat = "yozildi" if yozib_ol else "yangi etalon yaratildi"
            print(f"[{holat}] {nom}  ({joriy['review_count']} tekshirish bandi)")
            continue

        with open(yol, encoding="utf-8") as f:
            etalon = json.load(f)
        farq = _farqlar(etalon, joriy)
        if farq:
            xato += 1
            print(f"[FARQ] {nom}")
            print("\n".join(farq))
        else:
            print(f"[OK]   {nom}")

    if xato:
        print(f"\n{xato} ta faylda natija o'zgardi.")
        print("Agar o'zgarish ATAYLAB qilingan bo'lsa: python tests/golden.py --yozib-ol")
    return 1 if xato else 0


if __name__ == "__main__":
    sys.exit(main())
